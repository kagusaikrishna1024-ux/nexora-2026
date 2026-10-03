from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


OUTPUT = Path(__file__).with_name("NEXORA_2026_Presentation.pptx")
NAVY = RGBColor(18, 36, 58)
TEAL = RGBColor(0, 128, 128)
MINT = RGBColor(222, 242, 238)
ORANGE = RGBColor(228, 126, 55)
INK = RGBColor(35, 49, 63)
MUTED = RGBColor(91, 107, 122)
PAPER = RGBColor(247, 249, 250)
WHITE = RGBColor(255, 255, 255)
LINE = RGBColor(220, 227, 232)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def add_text(slide, x, y, width, height, value, size=18, color=INK,
             bold=False, font="Aptos", align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.03)
    frame.margin_right = Inches(0.03)
    frame.margin_top = Inches(0.02)
    frame.margin_bottom = Inches(0.02)
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_base(title, kicker="NEXORA 2026 | TECH INNOVATION HACKATHON"):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = PAPER
    band = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.12)
    )
    band.fill.solid()
    band.fill.fore_color.rgb = TEAL
    band.line.fill.background()
    add_text(slide, 0.7, 0.35, 9.5, 0.25, kicker, 10, TEAL, True)
    add_text(slide, 0.7, 0.7, 11.8, 0.62, title, 27, NAVY, True, "Aptos Display")
    footer = add_text(slide, 0.7, 7.14, 11.9, 0.2,
                      "Skillonex Academy  |  Student support and academy operations",
                      9, MUTED)
    footer.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
    return slide


def add_card(slide, x, y, width, height, title, body, accent=TEAL):
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x), Inches(y), Inches(width), Inches(height),
    )
    card.fill.solid()
    card.fill.fore_color.rgb = WHITE
    card.line.color.rgb = LINE
    stripe = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(0.08), Inches(height),
    )
    stripe.fill.solid()
    stripe.fill.fore_color.rgb = accent
    stripe.line.fill.background()
    add_text(slide, x + 0.25, y + 0.2, width - 0.45, 0.42,
             title, 17, NAVY, True, "Aptos Display")
    add_text(slide, x + 0.25, y + 0.72, width - 0.48, height - 0.88,
             body, 12, MUTED)


def add_step(slide, x, number, title, body):
    circle = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(x), Inches(2.0), Inches(0.55), Inches(0.55)
    )
    circle.fill.solid()
    circle.fill.fore_color.rgb = TEAL
    circle.line.fill.background()
    frame = circle.text_frame
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    run = frame.paragraphs[0].add_run()
    run.text = str(number)
    run.font.bold = True
    run.font.size = Pt(17)
    run.font.color.rgb = WHITE
    add_text(slide, x - 0.05, 2.72, 2.35, 0.4, title, 16, NAVY, True)
    add_text(slide, x - 0.05, 3.2, 2.35, 1.2, body, 12, MUTED)


# Cover
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = NAVY
accent = slide.shapes.add_shape(
    MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.22), prs.slide_height
)
accent.fill.solid()
accent.fill.fore_color.rgb = TEAL
accent.line.fill.background()
add_text(slide, 0.9, 0.75, 10.8, 0.35,
         "SKILLONEX ACADEMY  /  OFFICIAL HACKATHON PROTOTYPE", 12, MINT, True)
add_text(slide, 0.9, 1.6, 11.0, 1.2, "NEXORA 2026", 42, WHITE, True, "Aptos Display")
add_text(slide, 0.9, 2.8, 10.9, 1.25,
         "AI-powered student support, academy management, and complaint resolution",
         24, WHITE, True, "Aptos Display")
add_text(slide, 0.9, 4.65, 9.8, 0.55,
         "A unified student portal with explainable request triage and trackable support.",
         16, MINT)
add_text(slide, 0.9, 6.55, 9.0, 0.3,
         "TECH INNOVATION HACKATHON  |  3 OCTOBER 2026", 12, WHITE, True)

# Problem
slide = add_base("Students and support teams work across disconnected channels")
add_card(slide, 0.75, 1.65, 3.85, 3.7, "Student friction",
         "Course access, enrollment, payments, classes, certificates, and technical help are handled in separate places. Students cannot easily see what is happening next.", ORANGE)
add_card(slide, 4.75, 1.65, 3.85, 3.7, "Support workload",
         "Staff must read each request, determine the issue, judge urgency, choose a department, and follow progress manually.", TEAL)
add_card(slide, 8.75, 1.65, 3.85, 3.7, "Operational blind spots",
         "Without a shared queue and consistent issue data, managers lack a clear view of backlog, course demand, payment status, and recurring problems.", NAVY)
add_text(slide, 0.9, 5.85, 11.5, 0.65,
         "Challenge: understand the request, route it correctly, and help the academy respond faster.",
         18, NAVY, True)

# Solution
slide = add_base("One portal connects student services and academy support")
add_card(slide, 0.75, 1.65, 5.8, 2.0, "Student workspace",
         "Profile, courses, progress and attendance, class schedule, resources, payments, certificates, internships, announcements, and support history.")
add_card(slide, 6.8, 1.65, 5.8, 2.0, "Support assistant",
         "Natural-language request intake, topic classification, urgency detection, department routing, suggested reply, ticket ID, and status tracking.", ORANGE)
add_card(slide, 0.75, 3.95, 5.8, 2.0, "Academy operations",
         "Student and course overview, enrollment and payment signals, priority queue, resolution time, issue patterns, and management insights.", NAVY)
add_card(slide, 6.8, 3.95, 5.8, 2.0, "Working prototype",
         "Role-specific login, local SQLite persistence, seeded demo accounts, sample academy records, and admin actions that update the database.")

# Student journey
slide = add_base("Student journey: from enrollment to support resolution")
add_step(slide, 0.85, 1, "Sign in", "Open the student workspace and review personal course and service information.")
add_step(slide, 3.9, 2, "Learn", "Enroll in a course, track progress, and find class sessions and course resources.")
add_step(slide, 6.95, 3, "Ask", "Describe a payment, access, schedule, certificate, or other issue in natural language.")
add_step(slide, 10.0, 4, "Track", "See the triage result, suggested next step, ticket status, and support history.")
add_card(slide, 0.9, 5.0, 5.55, 1.2, "Personal services",
         "Payment records, certificates, internships, and academy updates are in the same portal.", TEAL)
add_card(slide, 6.8, 5.0, 5.55, 1.2, "Course enrollment",
         "Enrollment creates a course record and a pending payment entry for follow-up.", ORANGE)

# AI workflow
slide = add_base("Support triage turns free text into an actionable ticket")
steps = [
    (0.8, "Understand", "Normalize the student's text and compare intent terms."),
    (3.25, "Classify", "Assign a support category such as payment or course access."),
    (5.7, "Prioritize", "Detect explicit urgency and selected high-impact conditions."),
    (8.15, "Route", "Suggest the team that should handle the request."),
    (10.6, "Track", "Store a ticket and show its response and status to both roles."),
]
for number, (x, title, body) in enumerate(steps, start=1):
    add_step(slide, x, number, title, body)
add_card(slide, 0.85, 5.05, 5.65, 1.25, "Example: payment not reflected",
         "Payment / Enrollment  |  High  |  Accounts / Admissions", ORANGE)
add_card(slide, 6.8, 5.05, 5.65, 1.25, "Example: recording unavailable",
         "Course / Technical Support  |  Medium  |  Technical / Course Support", TEAL)

# Admin dashboard
slide = add_base("Admins see demand, priority, and service status")
metric_items = [
    (0.8, 1.65, "Student and course view", "30 demo students, enrollment totals, progress, and course-level activity."),
    (6.8, 1.65, "Support queue", "Open, in-progress, high-priority, and resolved requests with status controls."),
    (0.8, 3.75, "Payments", "Paid revenue and pending payment counts linked to course enrollments."),
    (6.8, 3.75, "Operational insight", "Average resolution time and most common request categories."),
]
for x, y, title, body in metric_items:
    add_card(slide, x, y, 5.7, 1.75, title, body)
add_text(slide, 0.9, 6.1, 11.5, 0.4,
         "Admin actions: update ticket status, add courses, and publish academy announcements.",
         15, NAVY, True)

# Architecture and AI
slide = add_base("Simple architecture keeps the prototype runnable and inspectable")
add_card(slide, 0.8, 1.65, 3.75, 2.1, "Streamlit UI",
         "Role-specific sign-in, student pages, support center, admin queue, and analytics.")
add_card(slide, 4.8, 1.65, 3.75, 2.1, "Python services",
         "Application workflows, local intent scoring, priority rules, suggested responses, and routing.", ORANGE)
add_card(slide, 8.8, 1.65, 3.75, 2.1, "SQLite data",
         "Users, courses, enrollments, payments, schedules, resources, certificates, announcements, and tickets.", NAVY)
add_text(slide, 0.95, 4.25, 11.1, 0.45, "AI approach and current limits", 19, NAVY, True)
add_text(slide, 0.95, 4.8, 11.2, 1.25,
         "The demo uses a deterministic local intent classifier rather than a hosted language model. It keeps student text on the local app, is easy to inspect, and handles the two stated examples. Ambiguous requests require staff review; multilingual understanding, sentiment, duplicate detection, and automatic escalation are future work.",
         15, MUTED)

# Technology stack
slide = add_base("Technology stack and implementation")
add_card(slide, 0.8, 1.65, 5.65, 1.75, "Application",
         "Python 3.13 with Streamlit for the interactive student and admin experience.", TEAL)
add_card(slide, 6.8, 1.65, 5.65, 1.75, "Data storage",
         "SQLite with parameterized queries for users, courses, enrollments, payments, and support tickets.", NAVY)
add_card(slide, 0.8, 3.75, 5.65, 1.75, "Security and configuration",
         "PBKDF2-HMAC-SHA256 password hashing with per-account salts; dark theme is configured in .streamlit/config.toml.", ORANGE)
add_card(slide, 6.8, 3.75, 5.65, 1.75, "AI implementation",
         "Python standard-library intent scoring and urgency rules provide local classification, response suggestions, and department routing.", TEAL)
add_text(slide, 0.9, 6.1, 11.5, 0.4,
         "No hosted model or external AI credential is required for the prototype.",
         15, NAVY, True)

# Demo and next steps
slide = add_base("Prototype results, demo access, and deployment next steps")
add_card(slide, 0.8, 1.65, 5.65, 2.0, "Demo access",
         "Student: student01 / Student@123\nKrishna: admin / Admin@123\nArsha: admin2 / Admin@123", TEAL)
add_card(slide, 6.8, 1.65, 5.65, 2.0, "Local verification",
         "30 seeded students, two admins, and six catalogue courses load into SQLite. The official triage examples and authenticated student/admin routes pass checks.", NAVY)
add_card(slide, 0.8, 4.0, 5.65, 1.8, "Deployment status",
         "The project currently runs locally. A public URL requires a connected source repository and deployment account; neither is configured in this workspace.", ORANGE)
add_card(slide, 6.8, 4.0, 5.65, 1.8, "Before real use",
         "Replace demo credentials, use persistent managed storage, add account recovery, and review AI suggestions before responding to students.", TEAL)

prs.save(OUTPUT)
print(f"Created {OUTPUT}")